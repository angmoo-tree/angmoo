"""Bounded upload validation and isolated document text extraction."""

from __future__ import annotations

from io import BytesIO
import multiprocessing
import os
from pathlib import PurePosixPath
import stat
import sys
from typing import Any, Callable
from zipfile import BadZipFile, ZipFile
from pypdf import PdfReader
from app.domains.character_lore.constants import (
    SUPPORTED_EXTENSIONS,
    MAX_LORE_FILE_BYTES,
    LORE_UPLOAD_READ_CHUNK_BYTES,
    DOCX_CONTENT_TYPE,
    MAX_DOCX_ENTRIES,
    MAX_DOCX_UNCOMPRESSED_BYTES,
    MAX_DOCX_XML_BYTES,
    MAX_DOCX_ENTRY_BYTES,
    MAX_DOCX_COMPRESSION_RATIO,
    MAX_PDF_PAGES,
    MAX_PDF_OBJECTS,
    MAX_PARSED_TEXT_CHARS,
    LORE_PARSER_TIMEOUT_SECONDS,
    LORE_PARSER_MEMORY_BYTES,
    LORE_PARSER_CPU_SECONDS,
)
from app.domains.character_lore.exceptions import (
    CharacterLoreValidationError,
    CharacterLoreFileTooLargeError,
)
from app.domains.character_lore.utils import _normalize_text


async def read_lore_upload_bytes(upload_file: Any) -> bytes:
    content = bytearray()
    while True:
        remaining = MAX_LORE_FILE_BYTES + 1 - len(content)
        if remaining <= 0:
            raise CharacterLoreFileTooLargeError(
                "설정집 파일은 10 MiB 이하만 업로드할 수 있습니다."
            )
        chunk = await upload_file.read(min(LORE_UPLOAD_READ_CHUNK_BYTES, remaining))
        if not chunk:
            return bytes(content)
        content.extend(chunk)
        if len(content) > MAX_LORE_FILE_BYTES:
            raise CharacterLoreFileTooLargeError(
                "설정집 파일은 10 MiB 이하만 업로드할 수 있습니다."
            )


def _validated_extension(filename: str) -> str:
    lower = filename.strip().lower()
    if lower.endswith(".doc") and not lower.endswith(".docx"):
        raise CharacterLoreValidationError("구형 .doc 파일은 v1에서 지원하지 않습니다.")
    extension = "." + lower.rsplit(".", 1)[-1] if "." in lower else ""
    if extension not in SUPPORTED_EXTENSIONS:
        raise CharacterLoreValidationError(
            "PDF, DOCX, TXT, MD 파일만 업로드할 수 있습니다."
        )
    return extension


def _safe_filename(filename: str) -> str:
    value = filename.strip().replace("\\", "/").split("/")[-1]
    return (value or "lore-file")[:240]


def validate_lore_upload_contract(
    *,
    extension: str,
    content_type: str | None,
    file_bytes: bytes,
) -> None:
    normalized_content_type = (content_type or "").split(";", 1)[0].strip().lower()
    allowed_content_types = {
        ".pdf": {"application/pdf"},
        ".docx": {DOCX_CONTENT_TYPE},
        ".txt": {"text/plain"},
        ".md": {"text/markdown", "text/plain", "text/x-markdown"},
    }
    if normalized_content_type not in allowed_content_types.get(extension, set()):
        raise CharacterLoreValidationError(
            "파일 확장자와 Content-Type이 일치하지 않습니다."
        )
    if not file_bytes:
        raise CharacterLoreValidationError("빈 설정집 파일은 업로드할 수 없습니다.")
    if extension == ".pdf":
        if not file_bytes.startswith(b"%PDF-"):
            raise CharacterLoreValidationError("유효한 PDF 파일이 아닙니다.")
        return
    if extension == ".docx":
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise CharacterLoreValidationError("유효한 DOCX 파일이 아닙니다.")
        _preflight_docx(file_bytes)
        return
    if b"\x00" in file_bytes or file_bytes.startswith((b"%PDF-", b"PK\x03\x04")):
        raise CharacterLoreValidationError(
            "텍스트 파일에서 바이너리 데이터가 감지됐습니다."
        )
    control_bytes = sum(
        byte < 32 and byte not in {9, 10, 13}
        for byte in file_bytes[: min(len(file_bytes), 64 * 1024)]
    )
    if control_bytes > max(1, min(len(file_bytes), 64 * 1024) // 100):
        raise CharacterLoreValidationError("유효한 텍스트 파일이 아닙니다.")


def _preflight_docx(file_bytes: bytes) -> None:
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise CharacterLoreValidationError("DOCX 내부 파일 수가 너무 많습니다.")
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise CharacterLoreValidationError("DOCX 필수 문서 구조가 없습니다.")

            total_uncompressed = 0
            total_xml = 0
            for entry in entries:
                path = PurePosixPath(entry.filename.replace("\\", "/"))
                if (
                    not path.parts
                    or path.is_absolute()
                    or ".." in path.parts
                    or ":" in path.parts[0]
                    or "\x00" in entry.filename
                ):
                    raise CharacterLoreValidationError(
                        "DOCX 내부 경로가 안전하지 않습니다."
                    )
                if entry.flag_bits & 0x1:
                    raise CharacterLoreValidationError(
                        "암호화된 DOCX 항목은 지원하지 않습니다."
                    )
                mode = entry.external_attr >> 16
                if stat.S_ISLNK(mode):
                    raise CharacterLoreValidationError(
                        "DOCX 내부 심볼릭 링크는 허용하지 않습니다."
                    )
                if entry.file_size > MAX_DOCX_ENTRY_BYTES:
                    raise CharacterLoreValidationError(
                        "DOCX 내부 파일 크기가 제한을 초과했습니다."
                    )
                if entry.file_size and (
                    entry.compress_size <= 0
                    or entry.file_size
                    > entry.compress_size * MAX_DOCX_COMPRESSION_RATIO
                ):
                    raise CharacterLoreValidationError(
                        "DOCX 압축 비율이 안전 한도를 초과했습니다."
                    )
                total_uncompressed += entry.file_size
                if total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise CharacterLoreValidationError(
                        "DOCX 압축 해제 크기가 제한을 초과했습니다."
                    )

                if entry.filename.lower().endswith((".xml", ".rels")):
                    total_xml += entry.file_size
                    if total_xml > MAX_DOCX_XML_BYTES:
                        raise CharacterLoreValidationError(
                            "DOCX XML 크기가 제한을 초과했습니다."
                        )
                    xml_bytes = archive.read(entry)
                    lowered = xml_bytes.lower()
                    if b"<!doctype" in lowered or b"<!entity" in lowered:
                        raise CharacterLoreValidationError(
                            "DOCX 외부 엔터티 선언은 허용하지 않습니다."
                        )
    except CharacterLoreValidationError:
        raise
    except (BadZipFile, OSError, RuntimeError, ValueError) as exc:
        raise CharacterLoreValidationError("유효한 DOCX 파일이 아닙니다.") from exc


def _extract_text(
    extension: str,
    file_bytes: bytes,
    *,
    content_type: str | None = None,
) -> str:
    if content_type is not None:
        validate_lore_upload_contract(
            extension=extension,
            content_type=content_type,
            file_bytes=file_bytes,
        )
    if extension in {".txt", ".md"}:
        text = _decode_text_file(file_bytes)
    elif extension == ".pdf":
        text = _extract_pdf_text(file_bytes)
    elif extension == ".docx":
        text = _extract_docx_text(file_bytes)
    else:
        raise CharacterLoreValidationError("지원하지 않는 설정집 파일 형식입니다.")
    normalized = _normalize_text(text)
    if not normalized:
        raise CharacterLoreValidationError("설정집에서 텍스트를 추출하지 못했습니다.")
    return normalized


def _decode_text_file(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        return file_bytes.decode("cp949", errors="replace")


def _extract_pdf_text(file_bytes: bytes) -> str:
    return _run_document_parser(".pdf", file_bytes)


def _extract_docx_text(file_bytes: bytes) -> str:
    return _run_document_parser(".docx", file_bytes)


def _run_document_parser(extension: str, file_bytes: bytes) -> str:
    result = _run_worker_process(
        _document_parser_worker,
        (extension, file_bytes),
        timeout_seconds=LORE_PARSER_TIMEOUT_SECONDS,
    )
    if not isinstance(result, str):
        raise CharacterLoreValidationError("문서 파서가 잘못된 결과를 반환했습니다.")
    return result


def _run_worker_process(
    target: Callable[..., None],
    args: tuple[Any, ...],
    *,
    timeout_seconds: float,
) -> Any:
    context = multiprocessing.get_context("spawn")
    parent_connection, child_connection = context.Pipe(duplex=False)
    process = context.Process(
        target=target,
        args=(child_connection, *args),
        daemon=True,
    )
    try:
        process.start()
        child_connection.close()
        if not parent_connection.poll(timeout_seconds):
            raise CharacterLoreValidationError(
                "문서 처리 시간이 안전 제한을 초과했습니다."
            )
        try:
            status_name, payload = parent_connection.recv()
        except (EOFError, OSError, ValueError) as exc:
            raise CharacterLoreValidationError(
                "격리된 문서 파서가 비정상 종료됐습니다."
            ) from exc
        if status_name == "ok":
            return payload
        if status_name == "validation":
            raise CharacterLoreValidationError(str(payload)[:500])
        raise CharacterLoreValidationError("문서 파서가 파일을 처리하지 못했습니다.")
    finally:
        parent_connection.close()
        child_connection.close()
        if process.is_alive():
            process.terminate()
            process.join(timeout=1)
        if process.is_alive():
            process.kill()
            process.join(timeout=1)
        if process.pid is not None and not process.is_alive():
            process.join(timeout=0.1)


def _document_parser_worker(
    connection: Any,
    extension: str,
    file_bytes: bytes,
) -> None:
    try:
        _apply_parser_resource_limits()
        if extension == ".pdf":
            text = _extract_pdf_text_in_process(file_bytes)
        elif extension == ".docx":
            text = _extract_docx_text_in_process(file_bytes)
        else:
            raise CharacterLoreValidationError("지원하지 않는 문서 형식입니다.")
        connection.send(("ok", text))
    except CharacterLoreValidationError as exc:
        connection.send(("validation", str(exc)[:500]))
    except BaseException:
        connection.send(("error", "document_parser_failed"))
    finally:
        connection.close()


def _extract_pdf_text_in_process(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(file_bytes))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise CharacterLoreValidationError("PDF 페이지 수가 제한을 초과했습니다.")
        object_count = reader.trailer.get("/Size", 0)
        if isinstance(object_count, int) and object_count > MAX_PDF_OBJECTS:
            raise CharacterLoreValidationError("PDF 객체 수가 제한을 초과했습니다.")
        parts: list[str] = []
        extracted_chars = 0
        for page in reader.pages:
            part = page.extract_text() or ""
            extracted_chars += len(part)
            if extracted_chars > MAX_PARSED_TEXT_CHARS:
                raise CharacterLoreValidationError(
                    "PDF 추출 텍스트가 제한을 초과했습니다."
                )
            if part.strip():
                parts.append(part)
    except CharacterLoreValidationError:
        raise
    except Exception as exc:
        raise CharacterLoreValidationError("PDF 텍스트를 추출하지 못했습니다.") from exc
    text = "\n\n".join(part.strip() for part in parts if part.strip())
    if not text.strip():
        raise CharacterLoreValidationError(
            "스캔 이미지 PDF/OCR은 v1에서 지원하지 않습니다."
        )
    return text


def _extract_docx_text_in_process(file_bytes: bytes) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(file_bytes))
    except Exception as exc:
        raise CharacterLoreValidationError(
            "DOCX 텍스트를 추출하지 못했습니다."
        ) from exc
    parts: list[str] = []
    extracted_chars = 0
    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        extracted_chars += len(paragraph.text)
        if extracted_chars > MAX_PARSED_TEXT_CHARS:
            raise CharacterLoreValidationError(
                "DOCX 추출 텍스트가 제한을 초과했습니다."
            )
        parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                row_text = " | ".join(cells)
                extracted_chars += len(row_text)
                if extracted_chars > MAX_PARSED_TEXT_CHARS:
                    raise CharacterLoreValidationError(
                        "DOCX 추출 텍스트가 제한을 초과했습니다."
                    )
                parts.append(row_text)
    return "\n\n".join(parts)


_WINDOWS_PARSER_JOB_HANDLE: int | None = None


def _apply_parser_resource_limits() -> None:
    if os.name == "nt":
        _apply_windows_parser_resource_limits()
        return
    if os.name == "posix":
        import resource

        resource.setrlimit(
            resource.RLIMIT_AS,
            (LORE_PARSER_MEMORY_BYTES, LORE_PARSER_MEMORY_BYTES),
        )
        resource.setrlimit(
            resource.RLIMIT_CPU,
            (LORE_PARSER_CPU_SECONDS, LORE_PARSER_CPU_SECONDS),
        )
        return
    raise RuntimeError(f"Unsupported parser isolation platform: {sys.platform}")


def _apply_windows_parser_resource_limits() -> None:
    import ctypes
    from ctypes import wintypes

    class IoCounters(ctypes.Structure):
        _fields_ = [
            ("ReadOperationCount", ctypes.c_ulonglong),
            ("WriteOperationCount", ctypes.c_ulonglong),
            ("OtherOperationCount", ctypes.c_ulonglong),
            ("ReadTransferCount", ctypes.c_ulonglong),
            ("WriteTransferCount", ctypes.c_ulonglong),
            ("OtherTransferCount", ctypes.c_ulonglong),
        ]

    class BasicLimitInformation(ctypes.Structure):
        _fields_ = [
            ("PerProcessUserTimeLimit", ctypes.c_longlong),
            ("PerJobUserTimeLimit", ctypes.c_longlong),
            ("LimitFlags", wintypes.DWORD),
            ("MinimumWorkingSetSize", ctypes.c_size_t),
            ("MaximumWorkingSetSize", ctypes.c_size_t),
            ("ActiveProcessLimit", wintypes.DWORD),
            ("Affinity", ctypes.c_size_t),
            ("PriorityClass", wintypes.DWORD),
            ("SchedulingClass", wintypes.DWORD),
        ]

    class ExtendedLimitInformation(ctypes.Structure):
        _fields_ = [
            ("BasicLimitInformation", BasicLimitInformation),
            ("IoInfo", IoCounters),
            ("ProcessMemoryLimit", ctypes.c_size_t),
            ("JobMemoryLimit", ctypes.c_size_t),
            ("PeakProcessMemoryUsed", ctypes.c_size_t),
            ("PeakJobMemoryUsed", ctypes.c_size_t),
        ]

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateJobObjectW.argtypes = [ctypes.c_void_p, wintypes.LPCWSTR]
    kernel32.CreateJobObjectW.restype = wintypes.HANDLE
    kernel32.SetInformationJobObject.argtypes = [
        wintypes.HANDLE,
        ctypes.c_int,
        ctypes.c_void_p,
        wintypes.DWORD,
    ]
    kernel32.SetInformationJobObject.restype = wintypes.BOOL
    kernel32.AssignProcessToJobObject.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
    kernel32.AssignProcessToJobObject.restype = wintypes.BOOL
    kernel32.GetCurrentProcess.restype = wintypes.HANDLE

    job = kernel32.CreateJobObjectW(None, None)
    if not job:
        raise OSError(ctypes.get_last_error(), "CreateJobObjectW failed")
    information = ExtendedLimitInformation()
    information.BasicLimitInformation.PerProcessUserTimeLimit = (
        LORE_PARSER_CPU_SECONDS * 10_000_000
    )
    information.BasicLimitInformation.LimitFlags = 0x00000002 | 0x00000100 | 0x00002000
    information.ProcessMemoryLimit = LORE_PARSER_MEMORY_BYTES
    if not kernel32.SetInformationJobObject(
        job,
        9,
        ctypes.byref(information),
        ctypes.sizeof(information),
    ):
        error_code = ctypes.get_last_error()
        kernel32.CloseHandle(job)
        raise OSError(error_code, "SetInformationJobObject failed")
    if not kernel32.AssignProcessToJobObject(job, kernel32.GetCurrentProcess()):
        error_code = ctypes.get_last_error()
        kernel32.CloseHandle(job)
        raise OSError(error_code, "AssignProcessToJobObject failed")
    global _WINDOWS_PARSER_JOB_HANDLE
    _WINDOWS_PARSER_JOB_HANDLE = int(job)
